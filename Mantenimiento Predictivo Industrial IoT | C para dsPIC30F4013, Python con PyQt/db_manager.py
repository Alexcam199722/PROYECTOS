"""
db_manager.py — Gestión de base de datos SQLite para el HMI de mantenimiento predictivo.

Tablas:
  sesiones        — una fila por sesión de monitoreo (arranque → detener)
  muestras        — serie temporal de cada sesión (c_proc, v_proc, rpm, di…)
  firma_arranque  — vector de firma de los primeros 3s de cada sesión
  eventos         — sobrecargas / desbalances con diagnóstico diferencial
  mantenimiento   — registro de intervenciones realizadas
"""

import sqlite3
import os
import json
import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    XLSX_OK = True
except ImportError:
    XLSX_OK = False


# ─────────────────────────────────────────────────────────────────────────────
DB_PATH = Path(__file__).parent / "motor_data.db"
# ─────────────────────────────────────────────────────────────────────────────


class DBManager:
    """Interfaz SQLite para el HMI de mantenimiento predictivo."""

    def __init__(self, db_path: str | Path = DB_PATH):
        self.db_path = Path(db_path)
        self._conn   = sqlite3.connect(str(self.db_path), check_same_thread=False)
        self._conn.row_factory = sqlite3.Row
        self._crear_tablas()
        self.sesion_id: int | None = None

    # ═════════════════════════════════════════════════════════════════════════
    # CREACIÓN DE TABLAS
    # ═════════════════════════════════════════════════════════════════════════
    def _crear_tablas(self):
        c = self._conn
        c.executescript("""
        PRAGMA journal_mode=WAL;

        CREATE TABLE IF NOT EXISTS sesiones (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            inicio          TEXT    NOT NULL,
            fin             TEXT,
            rpm_objetivo    REAL,
            sensor_tipo     TEXT,
            placa_json      TEXT,           -- datos de placa del fabricante
            muestras_total  INTEGER DEFAULT 0,
            horas_op        REAL    DEFAULT 0,
            sobrecargas     INTEGER DEFAULT 0,
            desbalances     INTEGER DEFAULT 0,
            di_medio        REAL,
            di_max          REAL,
            similitud_pct   REAL,           -- similitud con sesión sana de referencia
            notas           TEXT
        );

        CREATE TABLE IF NOT EXISTS muestras (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            sesion_id   INTEGER NOT NULL REFERENCES sesiones(id),
            t_muestra   INTEGER NOT NULL,   -- índice de muestra dentro de la sesión
            ts          TEXT    NOT NULL,   -- timestamp ISO
            c_proc      REAL,
            v_proc      REAL,
            rpm         REAL,
            c_crud      REAL,
            v_crud      REAL,
            ref_sana    REAL,
            bin_dom     INTEGER,
            di          REAL,
            predicho    INTEGER DEFAULT 0   -- 1 si el dato es predicho (microcorte)
        );

        CREATE TABLE IF NOT EXISTS firma_arranque (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            sesion_id   INTEGER NOT NULL REFERENCES sesiones(id),
            ts          TEXT    NOT NULL,
            vector_json TEXT    NOT NULL,   -- lista de dicts {t, c_proc, v_proc, rpm}
            rpm_pico    REAL,
            c_pico      REAL,
            t_estab_s   REAL,               -- tiempo hasta estabilización [s]
            similitud_pct REAL DEFAULT NULL -- similitud vs firma de referencia
        );

        CREATE TABLE IF NOT EXISTS eventos (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            sesion_id   INTEGER NOT NULL REFERENCES sesiones(id),
            ts          TEXT    NOT NULL,
            tipo        TEXT    NOT NULL,   -- SOBRECARGA / DESBALANCE
            hora_op     REAL,
            c_proc      REAL,
            c_crud      REAL,
            v_proc      REAL,
            v_crud      REAL,
            rpm         REAL,
            bin_dom     INTEGER,
            causa       TEXT,
            norma       TEXT,
            accion      TEXT
        );

        CREATE TABLE IF NOT EXISTS mantenimiento (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            ts          TEXT NOT NULL,
            componente  TEXT NOT NULL,
            sesion_id   INTEGER REFERENCES sesiones(id),
            horas_op    REAL,
            notas       TEXT
        );

        CREATE INDEX IF NOT EXISTS idx_muestras_sesion ON muestras(sesion_id);
        CREATE INDEX IF NOT EXISTS idx_eventos_sesion  ON eventos(sesion_id);
        """)
        c.commit()

    # ═════════════════════════════════════════════════════════════════════════
    # GESTIÓN DE SESIONES
    # ═════════════════════════════════════════════════════════════════════════
    def abrir_sesion(self, rpm_objetivo: float, sensor_tipo: str,
                     placa: dict | None = None) -> int:
        ts = datetime.datetime.now().isoformat(timespec='seconds')
        cur = self._conn.execute(
            "INSERT INTO sesiones (inicio, rpm_objetivo, sensor_tipo, placa_json) "
            "VALUES (?, ?, ?, ?)",
            (ts, rpm_objetivo, sensor_tipo, json.dumps(placa or {}))
        )
        self._conn.commit()
        self.sesion_id = cur.lastrowid
        return self.sesion_id

    def cerrar_sesion(self, horas_op: float, sobrecargas: int, desbalances: int,
                      di_medio: float | None, di_max: float | None,
                      similitud_pct: float | None, notas: str = ""):
        if self.sesion_id is None:
            return
        ts  = datetime.datetime.now().isoformat(timespec='seconds')
        cnt = self._conn.execute(
            "SELECT COUNT(*) FROM muestras WHERE sesion_id=?",
            (self.sesion_id,)
        ).fetchone()[0]
        self._conn.execute(
            "UPDATE sesiones SET fin=?, muestras_total=?, horas_op=?, "
            "sobrecargas=?, desbalances=?, di_medio=?, di_max=?, "
            "similitud_pct=?, notas=? WHERE id=?",
            (ts, cnt, horas_op, sobrecargas, desbalances,
             di_medio, di_max, similitud_pct, notas, self.sesion_id)
        )
        self._conn.commit()
        self.sesion_id = None

    # ═════════════════════════════════════════════════════════════════════════
    # ESCRITURA DE DATOS
    # ═════════════════════════════════════════════════════════════════════════
    def guardar_muestra(self, t_muestra: int, datos: dict, di: float | None = None,
                        predicho: bool = False):
        if self.sesion_id is None:
            return
        ts = datetime.datetime.now().isoformat(timespec='milliseconds')
        self._conn.execute(
            "INSERT INTO muestras (sesion_id,t_muestra,ts,c_proc,v_proc,rpm,"
            "c_crud,v_crud,ref_sana,bin_dom,di,predicho) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
            (self.sesion_id, t_muestra, ts,
             datos.get('c_proc'), datos.get('v_proc'), datos.get('rpm'),
             datos.get('c_crud'), datos.get('v_crud'), datos.get('ref_sana'),
             datos.get('bin_dom'), di, int(predicho))
        )
        # Commit cada 10 muestras para no saturar disco
        if t_muestra % 10 == 0:
            self._conn.commit()

    def guardar_evento(self, tipo: str, hora_op: float, datos: dict, diag: dict):
        if self.sesion_id is None:
            return
        ts = datetime.datetime.now().isoformat(timespec='seconds')
        self._conn.execute(
            "INSERT INTO eventos (sesion_id,ts,tipo,hora_op,c_proc,c_crud,"
            "v_proc,v_crud,rpm,bin_dom,causa,norma,accion) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (self.sesion_id, ts, tipo, hora_op,
             datos.get('c_proc'), datos.get('c_crud'),
             datos.get('v_proc'), datos.get('v_crud'),
             datos.get('rpm'),    datos.get('bin_dom'),
             diag.get('causa'),   diag.get('norma'), diag.get('accion'))
        )
        self._conn.commit()

    def guardar_firma(self, vector: list, rpm_pico: float, c_pico: float,
                      t_estab_s: float, similitud_pct: float | None = None):
        if self.sesion_id is None:
            return
        ts = datetime.datetime.now().isoformat(timespec='seconds')
        self._conn.execute(
            "INSERT INTO firma_arranque (sesion_id,ts,vector_json,rpm_pico,"
            "c_pico,t_estab_s,similitud_pct) VALUES (?,?,?,?,?,?,?)",
            (self.sesion_id, ts, json.dumps(vector),
             rpm_pico, c_pico, t_estab_s, similitud_pct)
        )
        self._conn.commit()

    # ═════════════════════════════════════════════════════════════════════════
    # LECTURA / ANÁLISIS ENTRE SESIONES
    # ═════════════════════════════════════════════════════════════════════════
    def obtener_sesiones(self, limite: int = 50) -> list[dict]:
        rows = self._conn.execute(
            "SELECT * FROM sesiones ORDER BY id DESC LIMIT ?", (limite,)
        ).fetchall()
        return [dict(r) for r in rows]

    def obtener_muestras_sesion(self, sesion_id: int) -> list[dict]:
        rows = self._conn.execute(
            "SELECT * FROM muestras WHERE sesion_id=? ORDER BY t_muestra",
            (sesion_id,)
        ).fetchall()
        return [dict(r) for r in rows]

    def obtener_firma_sesion(self, sesion_id: int) -> dict | None:
        row = self._conn.execute(
            "SELECT * FROM firma_arranque WHERE sesion_id=? ORDER BY id DESC LIMIT 1",
            (sesion_id,)
        ).fetchone()
        if row is None:
            return None
        d = dict(row)
        d['vector'] = json.loads(d['vector_json'])
        return d

    def obtener_firma_referencia(self) -> dict | None:
        """Firma de la sesión más antigua con DI bajo (motor sano)."""
        row = self._conn.execute(
            "SELECT s.id FROM sesiones s "
            "JOIN firma_arranque fa ON fa.sesion_id = s.id "
            "WHERE s.di_max IS NULL OR s.di_max < 1.5 "
            "ORDER BY s.id ASC LIMIT 1"
        ).fetchone()
        if row is None:
            return None
        return self.obtener_firma_sesion(row[0])

    def tendencia_di(self, n_sesiones: int = 20) -> list[dict]:
        """DI medio por sesión para el gráfico de tendencia histórica."""
        rows = self._conn.execute(
            "SELECT id, inicio, di_medio, di_max, sobrecargas, desbalances "
            "FROM sesiones WHERE di_medio IS NOT NULL "
            "ORDER BY id DESC LIMIT ?", (n_sesiones,)
        ).fetchall()
        return list(reversed([dict(r) for r in rows]))

    def similitud_historica(self, n_sesiones: int = 20) -> list[dict]:
        rows = self._conn.execute(
            "SELECT id, inicio, similitud_pct FROM sesiones "
            "WHERE similitud_pct IS NOT NULL ORDER BY id DESC LIMIT ?",
            (n_sesiones,)
        ).fetchall()
        return list(reversed([dict(r) for r in rows]))

    # ═════════════════════════════════════════════════════════════════════════
    # EXPORTACIÓN A EXCEL
    # ═════════════════════════════════════════════════════════════════════════
    def exportar_excel(self, ruta_out: str | Path | None = None) -> str:
        if not XLSX_OK:
            raise ImportError("openpyxl no instalado — ejecuta: pip install openpyxl")

        if ruta_out is None:
            ts  = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            ruta_out = self.db_path.parent / f"reporte_motor_{ts}.xlsx"

        wb = openpyxl.Workbook()

        # ── Estilos ────────────────────────────────────────────────────────
        DARK  = "0A0E1A"; BLUE  = "1A3060"; CYAN  = "00FFCC"
        GOLD  = "FFD700"; RED   = "FF3333"; ORANGE = "FF8C00"
        WHITE = "FFFFFF"; GRAY  = "556688"

        def hdr(ws, row, col, txt, bg=BLUE, fg=WHITE, bold=True):
            c = ws.cell(row, col, txt)
            c.fill = PatternFill("solid", fgColor=bg)
            c.font = Font(color=fg, bold=bold, name="Consolas", size=9)
            c.alignment = Alignment(horizontal="center", vertical="center")
            return c

        def cel(ws, row, col, val, fg=WHITE, fmt=None, bg=DARK):
            c = ws.cell(row, col, val)
            c.fill = PatternFill("solid", fgColor=bg)
            c.font = Font(color=fg, name="Consolas", size=8)
            c.alignment = Alignment(horizontal="center")
            if fmt:
                c.number_format = fmt
            return c

        thin = Border(
            left=Side(style='thin', color=BLUE),
            right=Side(style='thin', color=BLUE),
            top=Side(style='thin', color=BLUE),
            bottom=Side(style='thin', color=BLUE)
        )

        # ══════════════════════════════════════════════════════════════════
        # HOJA 1 — Resumen de sesiones
        # ══════════════════════════════════════════════════════════════════
        ws1 = wb.active
        ws1.title = "Resumen Sesiones"
        ws1.sheet_view.showGridLines = False
        ws1.column_dimensions['A'].width = 4
        for col, w in zip("BCDEFGHIJK", [22,10,8,10,10,10,10,10,14,14]):
            ws1.column_dimensions[col].width = w

        # Título
        ws1.merge_cells("B1:K1")
        t = ws1["B1"]
        t.value = "REPORTE HMI — MANTENIMIENTO PREDICTIVO DE MOTOR"
        t.fill  = PatternFill("solid", fgColor="002244")
        t.font  = Font(color=CYAN, bold=True, name="Consolas", size=12)
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws1.row_dimensions[1].height = 28

        ws1.merge_cells("B2:K2")
        sub = ws1["B2"]
        sub.value = f"Generado: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  DB: {self.db_path.name}"
        sub.fill  = PatternFill("solid", fgColor=DARK)
        sub.font  = Font(color=GRAY, name="Consolas", size=8)
        sub.alignment = Alignment(horizontal="center")

        hdrs1 = ["ID","Inicio","Horas Op","RPM Obj","DI Medio","DI Máx",
                 "Sobrecargas","Desbalances","Similitud %","Sensor"]
        for ci, h in enumerate(hdrs1, start=2):
            hdr(ws1, 4, ci, h)

        sesiones = self.obtener_sesiones(200)
        for ri, s in enumerate(sesiones, start=5):
            sim = s['similitud_pct']
            sim_col = CYAN if sim and sim > 80 else (GOLD if sim and sim > 60 else RED)
            di_col  = CYAN if (s['di_medio'] or 0) < 1.0 else (GOLD if (s['di_medio'] or 0) < 2.0 else RED)

            vals = [s['id'], s['inicio'], s['horas_op'] or 0,
                    s['rpm_objetivo'], s['di_medio'], s['di_max'],
                    s['sobrecargas'], s['desbalances'],
                    sim, s['sensor_tipo']]
            fgs  = [WHITE,WHITE,WHITE,WHITE, di_col, di_col,
                    (RED if (s['sobrecargas'] or 0) > 0 else WHITE),
                    (ORANGE if (s['desbalances'] or 0) > 0 else WHITE),
                    sim_col, WHITE]
            for ci, (v, fg) in enumerate(zip(vals, fgs), start=2):
                cel(ws1, ri, ci, v, fg=fg)

        # ══════════════════════════════════════════════════════════════════
        # HOJA 2 — Serie temporal de la última sesión
        # ══════════════════════════════════════════════════════════════════
        ws2 = wb.create_sheet("Última Sesión")
        ws2.sheet_view.showGridLines = False
        for col, w in zip("ABCDEFGHI", [5,22,10,10,10,10,10,6,10]):
            ws2.column_dimensions[col].width = w

        ws2.merge_cells("B1:I1")
        t2 = ws2["B1"]
        t2.value = "SERIE TEMPORAL — ÚLTIMA SESIÓN DE MONITOREO"
        t2.fill  = PatternFill("solid", fgColor="002244")
        t2.font  = Font(color=CYAN, bold=True, name="Consolas", size=11)
        t2.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 24

        hdrs2 = ["Timestamp","C_proc (mA)","V_proc","RPM","C_crudo (mA)","V_crudo","DI","Predicho"]
        for ci, h in enumerate(hdrs2, start=2):
            hdr(ws2, 3, ci, h)

        ultima = sesiones[0] if sesiones else None
        if ultima:
            muestras = self.obtener_muestras_sesion(ultima['id'])
            for ri, m in enumerate(muestras, start=4):
                di_col = CYAN if (m['di'] or 0) < 1.0 else (GOLD if (m['di'] or 0) < 2.0 else RED)
                vals = [m['ts'], m['c_proc'], m['v_proc'], m['rpm'],
                        m['c_crud'], m['v_crud'], m['di'], m['predicho']]
                fgs  = [WHITE,WHITE,WHITE,WHITE,WHITE,WHITE, di_col,
                        (GOLD if m['predicho'] else WHITE)]
                for ci, (v, fg) in enumerate(zip(vals, fgs), start=2):
                    cel(ws2, ri, ci, v, fg=fg)

            # Gráfico C_proc vs V_proc
            if len(muestras) > 2:
                chart = LineChart()
                chart.title = "Corriente EMA vs Vibración EMA"
                chart.style = 10
                chart.height = 12; chart.width  = 22
                n_rows = len(muestras)
                c_ref = Reference(ws2, min_col=3, min_row=4, max_row=3+n_rows)
                v_ref = Reference(ws2, min_col=4, min_row=4, max_row=3+n_rows)
                chart.add_data(c_ref); chart.add_data(v_ref)
                chart.series[0].title = SeriesLabel(v="C_proc (mA)")
                chart.series[1].title = SeriesLabel(v="V_proc")
                ws2.add_chart(chart, "B" + str(len(muestras) + 7))

        # ══════════════════════════════════════════════════════════════════
        # HOJA 3 — Eventos / Diagnósticos
        # ══════════════════════════════════════════════════════════════════
        ws3 = wb.create_sheet("Eventos y Diagnóstico")
        ws3.sheet_view.showGridLines = False
        for col, w in zip("ABCDEFGHIJ", [4,18,12,10,10,10,10,28,32,38]):
            ws3.column_dimensions[col].width = w

        ws3.merge_cells("B1:J1")
        t3 = ws3["B1"]
        t3.value = "HISTORIAL DE EVENTOS — DIAGNÓSTICO DIFERENCIAL"
        t3.fill  = PatternFill("solid", fgColor="330000")
        t3.font  = Font(color=RED, bold=True, name="Consolas", size=11)
        t3.alignment = Alignment(horizontal="center", vertical="center")
        ws3.row_dimensions[1].height = 24

        hdrs3 = ["Timestamp","Tipo","C_proc","C_crudo","V_proc","RPM","Causa probable","Norma ref.","Acción recomendada"]
        for ci, h in enumerate(hdrs3, start=2):
            hdr(ws3, 3, ci, h, bg="330000", fg=RED)

        eventos = self._conn.execute(
            "SELECT * FROM eventos ORDER BY id DESC LIMIT 500"
        ).fetchall()
        for ri, ev in enumerate(eventos, start=4):
            ev = dict(ev)
            color = "FF2222" if ev['tipo'] == 'SOBRECARGA' else "FF8800"
            vals  = [ev['ts'], ev['tipo'], ev['c_proc'], ev['c_crud'],
                     ev['v_proc'], ev['rpm'], ev['causa'], ev['norma'], ev['accion']]
            for ci, v in enumerate(vals, start=2):
                c = ws3.cell(ri, ci, v)
                c.fill = PatternFill("solid", fgColor="150000")
                c.font = Font(color=color, name="Consolas", size=8)
                c.alignment = Alignment(wrap_text=True, vertical="top")

        # ══════════════════════════════════════════════════════════════════
        # HOJA 4 — Tendencia histórica de DI
        # ══════════════════════════════════════════════════════════════════
        ws4 = wb.create_sheet("Tendencia DI")
        ws4.sheet_view.showGridLines = False
        for col, w in zip("ABCDEF", [4,22,12,12,12,14]):
            ws4.column_dimensions[col].width = w

        ws4.merge_cells("B1:F1")
        t4 = ws4["B1"]
        t4.value = "TENDENCIA HISTÓRICA — ÍNDICE DE DEGRADACIÓN (DI)"
        t4.fill  = PatternFill("solid", fgColor="001A33")
        t4.font  = Font(color=CYAN, bold=True, name="Consolas", size=11)
        t4.alignment = Alignment(horizontal="center", vertical="center")
        ws4.row_dimensions[1].height = 24

        hdrs4 = ["Inicio sesión","DI Medio","DI Máx","Sobrecargas","Similitud %"]
        for ci, h in enumerate(hdrs4, start=2):
            hdr(ws4, 3, ci, h, bg="001A33", fg=CYAN)

        tend = self.tendencia_di(100)
        for ri, t in enumerate(tend, start=4):
            di_m  = t['di_medio'] or 0
            di_col = CYAN if di_m < 1.0 else (GOLD if di_m < 2.0 else RED)
            vals   = [t['inicio'], t['di_medio'], t['di_max'],
                      t['sobrecargas'], None]
            sim_row = self._conn.execute(
                "SELECT similitud_pct FROM sesiones WHERE id=?", (t['id'],)
            ).fetchone()
            if sim_row:
                vals[4] = sim_row[0]

            for ci, v in enumerate(vals, start=2):
                cel(ws4, ri, ci, v, fg=di_col if ci in (3,4) else WHITE)

        # Gráfico DI histórico
        if len(tend) > 1:
            chart2 = LineChart()
            chart2.title = "Evolución DI entre sesiones"
            chart2.style = 10
            chart2.height = 12; chart2.width = 24
            di_ref = Reference(ws4, min_col=3, min_row=3, max_row=3+len(tend))
            chart2.add_data(di_ref, titles_from_data=True)
            ws4.add_chart(chart2, "B" + str(len(tend) + 7))

        wb.save(str(ruta_out))
        return str(ruta_out)

    def cerrar(self):
        self._conn.commit()
        self._conn.close()